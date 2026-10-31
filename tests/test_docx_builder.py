"""Kiem tra bao cao Word: thu tu muc, font, bang, anh, kiem cuoi. Du lieu gia, khong goi AI, khong goi R2.

Chay: python -m pytest tests/test_docx_builder.py -v
"""

from datetime import date

import docx
import pandas as pd
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

import config
from ledger import Ledger, Phase
from report.build import build_report
from report.charts import ChartError
from report.docx_builder import HEADINGS, ReportError, _shorten, build_document, report_title, save_document

TASK = "bai0a"
PASS = [{"name": "c", "status": "pass"}]
FONT = "Times New Roman"
PROBLEM = "# Bài 0a: Cơ cấu sản phẩm theo nhóm\n\nCâu hỏi: mỗi nhóm có bao nhiêu sản phẩm?"
CHART = {"type": "bar", "file": "by_group.csv", "x": "group", "y": ["total"], "series": None, "title": "Tổng theo nhóm"}
TABLE = {"file": "by_group.csv", "columns": ["group", "year", "total"]}
PROFILE = {"tables": {"PRODUCT": {"n_rows": 1250, "n_cols": 5, "flags": []}},
           "quality_flags": [{"severity": "info", "code": "x", "message": "Cột A thiếu một ít giá trị", "source_key": "k"}]}


def insight(i, requirement, **extra):
    return {"id": f"I{i}", "requirement": requirement, "text": "t", "text_filled": f"Nhận định {i} là 1.250.",
            "result_ids": ["R1"], "label": "verified", "table": None, "chart": None, **extra}


RECORD = {"summary": "s", "summary_filled": "Tóm tắt đạt 1.250.", "conclusion": "c", "conclusion_filled": "Kết luận rõ.",
          "limitations": "l", "limitations_filled": "Hạn chế rõ.", "requirements": ["Số sản phẩm", "Tổng giá"],
          "column_labels": {"group": "Nhóm hàng", "total": "Tổng tiền"}, "model": "fake",
          "insights": [insight(1, 1, table=TABLE, chart=CHART), insight(2, 2, label="hypothesis")]}


@pytest.fixture
def led(tmp_path):
    ledger = Ledger(tmp_path / "output" / TASK)
    ledger.set_state(Phase.DONE, plan={"requirements": RECORD["requirements"]})
    ledger.set_problem({"problem": PROBLEM, "restatement": {"restatement": "Đếm sản phẩm theo nhóm."}})
    work = ledger.root / "work"
    work.mkdir()
    pd.DataFrame({"group": [f"N{i}" for i in range(20)], "year": [2024] * 20,
                  "total": [1250.5 + i for i in range(20)]}).to_csv(work / "by_group.csv", index=False)
    ledger.add_result("Tổng theo nhóm?", "pandas groupby", {"n": 20}, PASS, question_key="Q1", output_path="work/by_group.csv")
    ledger.confirm_decision(ledger.add_decision("metric_definition", "Doanh thu gồm thuế").id)
    ledger.confirm_decision(ledger.add_decision("anomaly_handling", "Bỏ dòng trùng khóa").id)
    ledger.confirm_decision(ledger.add_decision("assumption", "Chấp nhận kết luận", source_key="accept_conclusion|synthesize").id)
    return ledger


def built(led, record=RECORD, mask=str):
    return docx.Document(build_report(led, TASK, record, PROFILE, mask)["docx"])


def styled(doc, name):
    return [p.text for p in doc.paragraphs if p.style.name == name]


def fonts(element):
    r = element.find(qn("w:rPr")).find(qn("w:rFonts"))
    return {k.split("}")[1]: v for k, v in r.attrib.items()}


# thu tu va noi dung
def test_seven_headings_in_order(led):
    assert styled(built(led), "Heading 1") == list(HEADINGS)


def test_requirements_are_subsections_in_order(led):
    h2 = styled(built(led), "Heading 2")
    assert h2[:2] == ["3.1. Số sản phẩm", "3.2. Tổng giá"]
    assert h2[2:] == ["5.1. Nguồn dữ liệu", "5.2. Chất lượng dữ liệu", "6.1. Phương pháp", "6.2. Giả định đã xác nhận"]


def test_cover_has_title_date_author_task_then_toc(led):
    doc = built(led)
    texts = [p.text for p in doc.paragraphs]
    assert styled(doc, "Title") == ["Cơ cấu sản phẩm theo nhóm"]
    assert any(t.startswith("Ngày lập: ") for t in texts) and "Người lập: " in texts and f"Tên bài: {TASK}" in texts
    assert texts.index(f"Tên bài: {TASK}") < texts.index("Mục lục") < texts.index(HEADINGS[0])


def test_toc_field_and_update_on_open(led):
    doc = built(led)
    assert any("TOC" in (i.text or "") for i in doc.element.body.iter(qn("w:instrText")))
    update = doc.settings.element.find(qn("w:updateFields"))
    assert update is not None and update.get(qn("w:val")) == "true"


def test_body_text_and_labels(led):
    doc = built(led)
    texts = [p.text for p in doc.paragraphs]
    assert "Nhận định 1 là 1.250." in texts and "Đã kiểm chứng" in texts and "Phỏng đoán" in texts
    assert "Tóm tắt đạt 1.250." in texts and "Kết luận rõ." in texts and "Hạn chế rõ." in texts
    run = next(p for p in doc.paragraphs if p.text == "Phỏng đoán").runs[0]
    assert run.font.size.pt == config.REPORT_LABEL_PT and run.font.color.rgb is not None


def test_captions_numbered(led):
    captions = styled(built(led), "Caption")
    assert captions == ["Bảng 1: Nhóm hàng, year, Tổng tiền (kết quả R1)", "Hình 1: Tổng theo nhóm"]


def test_sections_five_and_six_use_ledger(led):
    doc = built(led)
    bullets = styled(doc, "List Bullet")
    assert "Bỏ dòng trùng khóa" in bullets and "Cột A thiếu một ít giá trị" in bullets
    assert "Doanh thu gồm thuế" in bullets and "R1: Tổng theo nhóm? Phương pháp: pandas groupby" in bullets
    assert "Chấp nhận kết luận" not in bullets
    source = doc.tables[1]
    assert [c.text for c in source.rows[1].cells] == ["PRODUCT", "1.250", "5"]


def test_superseded_and_pending_decisions_left_out(led):
    led.add_decision("assumption", "CHUAXACNHAN")
    assert "CHUAXACNHAN" not in "\n".join(p.text for p in built(led).paragraphs)


def test_no_data_issue_says_so(led):
    led.supersede_decisions()
    doc = docx.Document(build_report(led, TASK, RECORD, {"tables": {}, "quality_flags": []})["docx"])
    texts = [p.text for p in doc.paragraphs]
    assert "Không phát hiện bất thường trong dữ liệu." in texts and "Không có giả định nào ngoài đề bài." in texts
    assert "Không có thông tin bảng nguồn." in texts


# dinh dang
@pytest.mark.parametrize("name", ["Normal", "Heading 1", "Heading 2", "Heading 3", "Caption", "List Bullet", "Title", "Table Grid"])
def test_times_new_roman_every_style_with_east_asian(led, name):
    got = fonts(built(led).styles[name].element)
    assert got == {"ascii": FONT, "hAnsi": FONT, "eastAsia": FONT, "cs": FONT}


def test_doc_defaults_font_and_body_size(led):
    doc = built(led)
    defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault"))
    assert fonts(defaults)["ascii"] == FONT and doc.styles["Normal"].font.size.pt == 12


def test_headings_keep_word_styles_black(led):
    style = built(led).styles["Heading 1"]
    color = style.element.find(qn("w:rPr")).find(qn("w:color"))
    assert color.get(qn("w:val")) == "000000" and color.get(qn("w:themeColor")) is None


def test_bullets_use_style_without_typed_character(led):
    doc = built(led)
    bullets = styled(doc, "List Bullet")
    assert bullets and not any(p.text.lstrip().startswith(("•", "-", "*")) for p in doc.paragraphs)


def test_page_setup_header_footer(led):
    section = built(led).sections[0]
    assert round(section.left_margin.cm, 1) == round(section.top_margin.cm, 1) == 2.5
    assert section.different_first_page_header_footer
    assert section.header.paragraphs[0].text == "Cơ cấu sản phẩm theo nhóm"
    assert section.first_page_header.paragraphs[0].text == ""
    codes = [i.text for i in section.footer._element.iter(qn("w:instrText"))]
    assert any("PAGE" in c for c in codes)


# bang
def test_table_header_style_and_repeat(led):
    table = built(led).tables[0]
    head = table.rows[0]
    assert head._tr.trPr.find(qn("w:tblHeader")) is not None
    assert all(c._tc.tcPr.find(qn("w:shd")) is not None for c in head.cells)
    assert all(c.paragraphs[0].runs[0].bold for c in head.cells)
    assert [c.text for c in head.cells] == ["Nhóm hàng", "year", "Tổng tiền"]
    assert table.style.name == "Table Grid"


def test_number_columns_right_aligned_vietnamese_year_unchanged(led):
    row = built(led).tables[0].rows[1]
    assert [c.text for c in row.cells] == ["N0", "2024", "1.250,5"]
    assert row.cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert row.cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert row.cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_long_table_cut_with_note(led):
    doc = built(led)
    assert len(doc.tables[0].rows) == 16
    assert any(p.text == "Hiển thị 15/20 dòng, đầy đủ trong file by_group.csv" for p in doc.paragraphs)


def test_short_table_no_note(led):
    pd.DataFrame({"group": ["A"], "year": [2024], "total": [1]}).to_csv(led.root / "work" / "by_group.csv", index=False)
    assert not any("Hiển thị" in p.text for p in built(led).paragraphs)


# anh va kiem cuoi
def test_picture_embedded(led):
    assert len(built(led).inline_shapes) == 1


def test_no_placeholder_left_in_good_report(led):
    doc = built(led)
    assert "{{" not in "\n".join(p.text for p in doc.paragraphs)


@pytest.mark.parametrize("field", ["summary_filled", "conclusion_filled", "limitations_filled"])
def test_placeholder_left_stops_without_file(led, field):
    with pytest.raises(ReportError, match="placeholder"):
        build_report(led, TASK, {**RECORD, field: "còn {{R1.n}}"}, PROFILE)
    assert not list((led.root / "report").glob("*.docx"))


def test_missing_table_stops_without_file(led):
    record = {**RECORD, "insights": [insight(1, 1, table={"file": "khong_co.csv", "columns": ["a"]}), insight(2, 2)]}
    with pytest.raises(ReportError, match="không có bảng khong_co.csv"):
        build_report(led, TASK, record, PROFILE)
    assert not list((led.root / "report").glob("*.docx"))


def test_missing_column_stops(led):
    record = {**RECORD, "insights": [insight(1, 1, table={"file": "by_group.csv", "columns": ["khong_co"]}), insight(2, 2)]}
    with pytest.raises(ReportError, match="không có cột khong_co"):
        build_report(led, TASK, record, PROFILE)


def test_cover_date_format(led):
    doc = build_document({**RECORD, "insights": []}, led, {}, TASK, PROFILE, today=date(2026, 10, 31))
    assert "Ngày lập: 31/10/2026" in [p.text for p in doc.paragraphs]


def test_missing_image_stops(led):
    with pytest.raises(ReportError, match="không có ảnh biểu đồ"):
        build_document(RECORD, led, {}, TASK, PROFILE)


def test_missing_chart_file_stops_before_word(led):
    (led.root / "work" / "by_group.csv").unlink()
    with pytest.raises(ChartError):
        build_report(led, TASK, RECORD, PROFILE)


def test_failed_rebuild_keeps_previous_file(led):
    path = build_report(led, TASK, RECORD, PROFILE)["docx"]
    before = path.read_bytes()
    with pytest.raises(ReportError):
        build_report(led, TASK, {**RECORD, "summary_filled": "{{R1.n}}"}, PROFILE)
    assert path.read_bytes() == before


def test_open_file_gives_clear_error(led, monkeypatch):
    doc = built(led)

    def denied(*a):
        raise PermissionError

    monkeypatch.setattr("os.replace", denied)
    with pytest.raises(ReportError, match="đang mở trong Word"):
        save_document(doc, led.root / "report" / "x.docx")
    assert not list((led.root / "report").glob(".*.tmp"))


def test_masked_value_never_in_word(led):
    pd.DataFrame({"group": ["BIMAT", "A"], "year": [2024, 2024], "total": [1, 2]}).to_csv(
        led.root / "work" / "by_group.csv", index=False)
    record = {**RECORD, "conclusion_filled": "Khách BIMAT mua nhiều."}
    doc = built(led, record, mask=lambda s: s.replace("BIMAT", "[đã che]"))
    text = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    assert "BIMAT" not in text and "[đã che]" in text


def test_masked_numeric_id_not_shown_as_number(led):
    pd.DataFrame({"group": ["A"], "year": [2024], "total": [778899]}).to_csv(led.root / "work" / "by_group.csv", index=False)
    doc = built(led, {**RECORD, "insights": [insight(1, 1, table=TABLE), insight(2, 2)]},
                mask=lambda s: s.replace("778899", "[đã che]"))
    assert doc.tables[0].rows[1].cells[2].text == "[đã che]"


def test_old_record_without_new_keys(led):
    old = {k: v for k, v in RECORD.items() if k != "column_labels"}
    old["insights"] = [{k: v for k, v in insight(i, i).items() if k not in ("table", "chart")} for i in (1, 2)]
    assert styled(built(led, old), "Heading 1") == list(HEADINGS)


# phu tro
@pytest.mark.parametrize("problem, title", [("# Bài 2: Đánh giá hiệu quả\n\nx", "Đánh giá hiệu quả"), ("\n# Đề khác\n", "Đề khác"),
                                            ("", "bai0a")])
def test_report_title(problem, title):
    assert report_title({"problem": problem}, "bai0a") == title


def test_summary_cut_at_sentence():
    text = "Câu một. " * 400
    short = _shorten(text, config.REPORT_SUMMARY_MAX_CHARS)
    assert len(short) <= config.REPORT_SUMMARY_MAX_CHARS and short.endswith(".")
    assert _shorten("Ngắn.", 100) == "Ngắn."
