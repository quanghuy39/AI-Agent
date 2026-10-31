"""Bao cao Word gui lanh dao, dung tu insights.json. Chi code, khong AI: moi so da dien san hoac doc tu file bang.

Moi chuoi dua vao file deu qua mask: text_filled dien gia tri that nen co the chua gia tri cot che, du AI khong thay no.
Co loi (anh, bang thieu, con placeholder) thi ReportError, khong ghi file; ghi ra file tam roi moi thay file that.
"""

import os
import re
import tempfile
from datetime import date

import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pandas.api.types import is_bool_dtype, is_numeric_dtype

import config
from analysis.loop import ACCEPT_FLAGS_KEY
from ledger import DecisionStatus, ResultStatus
from profiling import flags as fl
from profiling import quality
from report.charts import _label, tables_of
from report.placeholders import HYPOTHESIS, VERIFIED, format_value
from report.tables import load_table

# decision chap nhan ket luan: co dinh, moi lan tu choi roi tong hop lai dung chung; khong phai gia dinh phan tich
CONCLUSION_KEY = "accept_conclusion|synthesize"
ANOMALY = "anomaly_handling"
LABELS = {VERIFIED: "Đã kiểm chứng", HYPOTHESIS: "Phỏng đoán"}
GRAY = RGBColor(0x6B, 0x6B, 0x6B)
BLACK = RGBColor(0, 0, 0)
HEADINGS = ("1. Tóm tắt", "2. Mục tiêu", "3. Kết quả", "4. Kết luận", "5. Dữ liệu và chất lượng",
            "6. Phương pháp và giả định", "7. Hạn chế")
HEADING_SIZES = {"Title": 24, "Heading 1": 16, "Heading 2": 14, "Heading 3": 13}


class ReportError(Exception):
    """Khong dung duoc bao cao hop le; message liet ke loi."""


# ---------- dinh dang chung ----------

def _set_fonts(rpr):
    """Mot font cho moi bang ma, bo thuoc tinh theme (theme se de len font khai bao)."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.attrib.clear()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), config.REPORT_FONT)


def _black(style):
    """Bo mau theme cua style co san trong mau Word."""
    rpr = style.element.get_or_add_rPr()
    for color in rpr.findall(qn("w:color")):
        rpr.remove(color)
    style.font.color.rgb = BLACK


def _setup_styles(doc):
    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.LIST:
            _set_fonts(style.element.get_or_add_rPr())
    defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _set_fonts(defaults)
    for lang in defaults.findall(qn("w:lang")):
        lang.set(qn("w:val"), "vi-VN")
    normal = doc.styles["Normal"]
    normal.font.size = Pt(config.REPORT_BODY_PT)
    normal.paragraph_format.space_after = Pt(6)
    for name, size in HEADING_SIZES.items():
        doc.styles[name].font.size = Pt(size)
        _black(doc.styles[name])
    caption = doc.styles["Caption"]
    caption.font.size = Pt(config.REPORT_TABLE_PT)
    _black(caption)
    _black(doc.styles["TOC Heading"])


def _setup_page(doc, title):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    margin = Cm(config.REPORT_MARGIN_CM)
    section.left_margin = section.right_margin = section.top_margin = section.bottom_margin = margin
    section.different_first_page_header_footer = True  # trang bia khong co dau trang, so trang
    section.header.paragraphs[0].text = title
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_field(footer, "PAGE", "1")


def _add_field(paragraph, instruction, placeholder):
    """Truong cua Word (PAGE, TOC): gia tri that do Word tinh khi mo file."""
    def char(kind):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        return element

    run = paragraph.add_run()
    run._r.append(char("begin"))
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = f" {instruction} "
    run = paragraph.add_run()
    run._r.append(code)
    paragraph.add_run()._r.append(char("separate"))
    paragraph.add_run(placeholder)
    paragraph.add_run()._r.append(char("end"))


def _update_fields_on_open(doc):
    element = OxmlElement("w:updateFields")
    element.set(qn("w:val"), "true")
    doc.settings.element.append(element)


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------- bang ----------

def _shade(cell):
    fill = OxmlElement("w:shd")
    fill.set(qn("w:val"), "clear")
    fill.set(qn("w:color"), "auto")
    fill.set(qn("w:fill"), config.REPORT_HEADER_FILL)
    cell._tc.get_or_add_tcPr().append(fill)


def _row_flag(row, name):
    row._tr.get_or_add_trPr().append(OxmlElement(name))


def _write_cell(cell, text, right, bold=False):
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if right else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(config.REPORT_TABLE_PT)


def _add_table(doc, headers, rows, right=()):
    """Bang vien manh, dong tieu de to nen, chu dam, lap lai khi sang trang; right: chi so cot can phai."""
    table = doc.add_table(rows=1, cols=len(headers), style="Table Grid")
    _row_flag(table.rows[0], "w:tblHeader")
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        _shade(cell)
        _write_cell(cell, text, i in right, bold=True)
    for values in rows:
        row = table.add_row()
        _row_flag(row, "w:cantSplit")
        for i, (cell, text) in enumerate(zip(row.cells, values)):
            _write_cell(cell, text, i in right)
    return table


def _is_measure(name, column):
    """Cot so do: can phai, dinh dang so. Cot so mang nghia ma (nam, id) giu nguyen de khong in 2024 thanh 2.024."""
    if not is_numeric_dtype(column) or is_bool_dtype(column):
        return False
    words = {w.lower() for w in quality.name_words(str(name))}
    return not words & set(config.REPORT_PLAIN_NUMBER_HINTS)


def _cell_text(value, measure, mask):
    if pd.isna(value):
        return ""
    plain = _label(value.item() if hasattr(value, "item") else value)
    masked = mask(plain)
    if masked != plain or not measure:
        return masked
    number = float(value)
    return format_value(int(number) if number.is_integer() else number)


# ---------- noi dung ----------

class _Report:
    """Trang thai khi dung: danh so bang, hinh va loi gom lai."""

    def __init__(self, doc, record, ledger, charts, mask):
        self.doc, self.record, self.ledger, self.charts, self.mask = doc, record, ledger, charts, mask
        self.labels = record.get("column_labels") or {}
        self.n_tables = self.n_figures = 0
        self.errors = []

    def text(self, value, style=None):
        return self.doc.add_paragraph(self.mask(value), style=style)

    def bullets(self, items):
        for item in items:
            self.text(item, "List Bullet")

    def heading(self, value, level):
        self.doc.add_heading(self.mask(value), level)

    def _caption(self, value, keep_with_next):
        paragraph = self.text(value, "Caption")
        paragraph.paragraph_format.keep_with_next = keep_with_next

    def label(self, label):
        run = self.doc.add_paragraph().add_run(LABELS[label])
        run.font.size = Pt(config.REPORT_LABEL_PT)
        run.font.color.rgb = GRAY

    def data_table(self, ins):
        spec = ins["table"]
        table = tables_of(self.ledger, ins["result_ids"]).get(spec["file"])
        if table is None:
            self.errors.append(f"{ins['id']}: không có bảng {spec['file']}.")
            return
        df = load_table(table["path"])
        missing = [c for c in spec["columns"] if c not in df.columns]
        if missing:
            self.errors.append(f"{ins['id']}: bảng {spec['file']} không có cột {', '.join(missing)}.")
            return
        self.n_tables += 1
        headers = [self.labels.get(c, c) for c in spec["columns"]]
        cited = ", ".join(ins["result_ids"])
        self._caption(f"Bảng {self.n_tables}: {', '.join(headers)} (kết quả {cited})", True)
        measures = [_is_measure(c, df[c]) for c in spec["columns"]]
        top = df.head(config.REPORT_TABLE_MAX_ROWS)
        rows = [[_cell_text(top[c].iloc[i], m, self.mask) for c, m in zip(spec["columns"], measures)]
                for i in range(len(top))]
        _add_table(self.doc, [self.mask(h) for h in headers], rows, {i for i, m in enumerate(measures) if m})
        if len(df) > len(top):
            self._small(f"Hiển thị {len(top)}/{len(df)} dòng, đầy đủ trong file {spec['file']}")
        else:
            self.doc.add_paragraph()

    def _small(self, value):
        run = self.doc.add_paragraph().add_run(self.mask(value))
        run.font.size = Pt(config.REPORT_LABEL_PT)
        run.italic = True

    def figure(self, ins):
        path = self.charts.get(ins["id"])
        if path is None or not path.is_file():
            self.errors.append(f"{ins['id']}: không có ảnh biểu đồ.")
            return
        self.n_figures += 1
        paragraph = self.doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.keep_with_next = True
        paragraph.add_run().add_picture(str(path), width=Cm(config.REPORT_PICTURE_CM))
        self._caption(f"Hình {self.n_figures}: {ins['chart']['title']}", False)


def report_title(problem, fallback):
    """Dong tieu de dau tien cua de, bo dau # va tien to "Bài ...:"."""
    for line in (problem.get("problem") or "").splitlines():
        line = re.sub(r"^#+\s*", "", line.strip())
        if line:
            return re.sub(r"^Bài\s+\S+\s*:\s*", "", line, flags=re.I) or line
    return fallback


def _shorten(text, limit):
    """Cat o cuoi cau gan nhat, de muc tom tat khong qua nua trang."""
    if len(text) <= limit:
        return text
    cut = text[:limit]
    end = cut.rfind(". ")
    return cut[:end + 1] if end > 0 else cut.rstrip() + "…"


def _cover(doc, title, mask, today):
    for _ in range(5):
        doc.add_paragraph()
    doc.add_paragraph(mask(title), style="Title")
    doc.add_paragraph(f"Ngày lập: {today.strftime('%d/%m/%Y')}")
    author = config.report_author()
    if author:
        doc.add_paragraph(f"Người lập: {author}")
    _page_break(doc)


def _contents(doc):
    doc.add_paragraph("Mục lục", style="TOC Heading")
    _add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u', "Chọn cập nhật khi Word hỏi để hiện mục lục.")
    _page_break(doc)


def _results(rep):
    record = rep.record
    for i, requirement in enumerate(record["requirements"], 1):
        rep.heading(f"3.{i}. {requirement}", 2)
        for ins in (x for x in record["insights"] if x["requirement"] == i):
            rep.text(ins["text_filled"])
            rep.label(ins["label"])
            if ins.get("table"):
                rep.data_table(ins)
            if ins.get("chart"):
                rep.figure(ins)


def _names_result(decision):
    """Decision "chap nhan R.. du co co": co ma ket qua nen chi hien o muc 6, muc giu ma de truy vet."""
    return (decision.source_key or "").startswith(ACCEPT_FLAGS_KEY)


def _data_section(rep, profile):
    rep.heading("5.1. Nguồn dữ liệu", 2)
    rows = [[name, format_value(t["n_rows"]), format_value(t["n_cols"])] for name, t in (profile.get("tables") or {}).items()]
    if rows:
        _add_table(rep.doc, ["Bảng nguồn", "Số dòng", "Số cột"], [[rep.mask(c) for c in r] for r in rows], {1, 2})
        rep.doc.add_paragraph()
    else:
        rep.text("Không có thông tin bảng nguồn.")
    rep.heading("5.2. Chất lượng dữ liệu", 2)
    handled = [d.text for d in rep.ledger.active_decisions()
               if d.status == DecisionStatus.CONFIRMED and d.kind.value == ANOMALY and not _names_result(d)]
    notes = [f["message"] for f in fl.all_flags(profile) if f["severity"] == fl.INFO]
    if handled or notes:
        rep.bullets(handled + notes)
    else:
        rep.text("Không phát hiện bất thường trong dữ liệu.")


def _method_section(rep):
    rep.heading("6.1. Phương pháp", 2)
    valid = [r for r in rep.ledger.results() if r.status == ResultStatus.VALID]
    rep.bullets(f"{r.id}: {r.question} Phương pháp: {r.method}" for r in valid)
    rep.heading("6.2. Giả định đã xác nhận", 2)
    assumptions = [d.text for d in rep.ledger.active_decisions()
                   if d.status == DecisionStatus.CONFIRMED and d.source_key != CONCLUSION_KEY
                   and (d.kind.value != ANOMALY or _names_result(d))]
    if assumptions:
        rep.bullets(assumptions)
    else:
        rep.text("Không có giả định nào ngoài đề bài.")


def _all_text(doc):
    parts = [p.text for p in doc.paragraphs] + [p.text for p in doc.sections[0].header.paragraphs]
    parts += [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    return "\n".join(parts)


def build_document(record, ledger, charts, task, profile, mask=str, today=None):
    """Dung Document trong bo nho theo C1. Co loi -> ReportError."""
    problem = ledger.get_problem() or {}
    title = report_title(problem, task)
    doc = Document()
    doc.core_properties.title = mask(title)
    doc.core_properties.author = ""
    _setup_styles(doc)
    _setup_page(doc, mask(title))
    _update_fields_on_open(doc)
    _cover(doc, title, mask, today or date.today())
    _contents(doc)

    rep = _Report(doc, record, ledger, charts, mask)
    rep.heading(HEADINGS[0], 1)
    rep.text(_shorten(record["summary_filled"], config.REPORT_SUMMARY_MAX_CHARS))
    rep.heading(HEADINGS[1], 1)
    restatement = (problem.get("restatement") or {}).get("restatement")
    if restatement:
        rep.text(restatement)
    rep.bullets(record["requirements"])
    rep.heading(HEADINGS[2], 1)
    _results(rep)
    rep.heading(HEADINGS[3], 1)
    rep.text(record["conclusion_filled"])
    rep.heading(HEADINGS[4], 1)
    _data_section(rep, profile)
    rep.heading(HEADINGS[5], 1)
    _method_section(rep)
    rep.heading(HEADINGS[6], 1)
    rep.text(record["limitations_filled"])

    if "{{" in _all_text(doc):
        rep.errors.append("Báo cáo còn placeholder chưa điền.")
    if rep.errors:
        raise ReportError(" ".join(rep.errors))
    return doc


def save_document(doc, path):
    """Ghi file tam roi thay file that: loi giua chung khong de lai file do dang."""
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(dir=path.parent, prefix=f".{path.name}.", suffix=".tmp")
    os.close(fd)
    try:
        doc.save(tmp)
        os.replace(tmp, path)
    except PermissionError:
        raise ReportError(f"Không ghi được {path.name}: file đang mở trong Word, đóng file rồi chạy lại.") from None
    finally:
        if os.path.exists(tmp):
            os.remove(tmp)
